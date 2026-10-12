"""Customer-facing workbook contract, using synthetic evidence."""
import copy
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from build_report import build, normalize, complete_for_export
from test_report import fixture


class ReportFormatTests(unittest.TestCase):
    def test_defect_and_evidence_gap_keep_both_actions(self):
        data=fixture()
        data['evidence_gaps']=[{'check':'10.1','missing':'Historical response missing.','next_action':'Provide the earlier response.'}]
        rows,_=normalize(data)
        self.assertIn(data['issues'][0]['how_to_fix'],rows[0]['instruction'])
        self.assertIn('Provide the earlier response.',rows[0]['instruction'])
        self.assertNotIn('Provide the earlier response.',rows[0]['findings'])

    def test_new_columns_symbols_and_numbered_shared_issue(self):
        data=fixture()
        data['checklist'][0]['instruction']='Add the missing rules.'
        data['checklist'][1]['result']='No'
        data['issues'][0]['check_ids'].append('10.4')
        for row in data['checklist']:
            row.pop('coverage')
        with tempfile.TemporaryDirectory(dir=Path.cwd()) as directory:
            output=Path(directory)/'report.xlsx'
            build(data,output)
            wb=load_workbook(output)
            self.assertEqual([c.value for c in wb['Checklist'][1]],['Check','Result','Findings','Instruction'])
            self.assertEqual(wb['Checklist']['B2'].value,'X')
            self.assertEqual(wb['Checklist']['D2'].value,'Add the missing rules.')
            self.assertEqual(wb['10. Robot.txt']['A2'].value,'10.1, 10.4 Empty robots.txt')
            wb.close()

    def test_human_action_moves_out_of_findings_and_no_review_issue(self):
        data={'requested_checks':['10.4'],'checklist':[{
            'id':'10.4','check':'old title','result':'Human Check',
            'findings':'Public file usage is unknown.\nNext action: Export the file links.',
            'human_check_reason':'Public file usage is unknown.',
            'human_check_action':'Export the file links.'}], 'issues':[]}
        rows,issues=normalize(data)
        self.assertNotIn('Export the file links.',rows[0]['findings'])
        self.assertEqual(rows[0]['instruction'],'Export the file links.')
        self.assertEqual(issues,[])
        self.assertEqual(normalize({'requested_checks':['10.4'],'checklist':rows,'issues':[]})[0],rows)

    def test_symbol_input_and_instruction_formula_safety(self):
        data=fixture();data['checklist'][0]['result']='X';data['checklist'][1]['result']='√'
        data['checklist'][0]['instruction']='=untrusted instruction'
        with tempfile.TemporaryDirectory(dir=Path.cwd()) as directory:
            output=Path(directory)/'report.xlsx';build(data,output)
            wb=load_workbook(output)
            self.assertEqual(wb['Checklist']['B3'].value,'√')
            self.assertEqual(wb['Checklist']['D2'].data_type,'s')
            self.assertEqual(wb['Checklist']['D2'].value,'=untrusted instruction')
            self.assertEqual(wb['Checklist']['D3'].value,'—')
            wb.close()

    def test_symbol_pass_with_gap_is_not_exported_as_pass(self):
        data=fixture();data['checklist'][1]['result']='√'
        data['evidence_gaps']=[{'check':'10.4','missing':'File links missing.','next_action':'Export file links.'}]
        prepared=complete_for_export(data,'https://example.com/')
        self.assertEqual(prepared['checklist'][1]['result'],'Human Check')
        self.assertIn('Export file links.',prepared['checklist'][1]['instruction'])

if __name__=='__main__': unittest.main()
