"""Offline contract checks using synthetic data, not customer crawl results."""
import copy
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from build_report import build, complete_for_export, normalize, report_filename


def fixture():
    return {'requested_checks':['10.1','10.4','10.5'],'checklist':[
        {'id':'10.1','check':'10.1 Does robots.txt exist?','result':'No','findings':'Empty robots.txt observed.','coverage':'One synthetic response.'},
        {'id':'10.4','check':'10.4 CSS/JS access','result':'Yes','findings':'Required resources allowed and fetched.','coverage':'Two synthetic CSS/JS resources.'},
        {'id':'10.5','check':'10.5 Special pages','result':'NA','findings':'No verifiable objects.','na_reason':'15 candidate paths tested; none verifiable; no transaction performed.','coverage':'Synthetic candidates.'}],
        'issues':[{'check_ids':['10.1'],'issue':'Empty robots.txt','description':'=untrusted description','how_to_fix':'Configure appropriate directives.','addresses':['https://example.com/robots.txt?x=1&y=2']}]}


class ReportTests(unittest.TestCase):
    def test_partial_evidence_survives_in_checklist_and_generated_review_issue(self):
        data={'requested_checks':['10.2'],'checklist':[{
            'id':'10.2','check':'10.2 Important URLs','result':'Human Check',
            'findings':'Read robots and checked 75 sitemap URLs: allowed.',
            'coverage':'75/80 permissions checked; five unresolved.'}],
            'issues':[], 'evidence_gaps':[{
                'check':'10.2','missing':'Five effective permission results unavailable after SF export error.',
                'next_action':'Obtain the five unresolved permission results.'}]}
        prepared=complete_for_export(data,'https://example.com/')
        with tempfile.TemporaryDirectory(dir=Path.cwd()) as directory:
            output=Path(directory)/'audit.xlsx';build(prepared,output)
            wb=load_workbook(output)
            self.assertIn('75 sitemap URLs',wb['Checklist']['C2'].value)
            self.assertIn('five unresolved',wb['Checklist']['D2'].value)
            self.assertIn('75 sitemap URLs',wb['Issues']['B2'].value)
            self.assertIn('Five effective permission results unavailable',wb['Issues']['B2'].value)
            self.assertIn('five unresolved permission results',wb['Issues']['C2'].value)
            wb.close()

    def test_delivery_completion_exports_missing_rows_without_inventing_passes(self):
        data=fixture(); del data['requested_checks']; data['checklist'].pop(1)
        original=copy.deepcopy(data)
        prepared=complete_for_export(data, 'https://audit-client.test/path')
        self.assertEqual(data,original)
        self.assertEqual(len(prepared['checklist']),7)
        results={row['id']:row['result'] for row in prepared['checklist']}
        self.assertEqual(results['10.1'],'No'); self.assertEqual(results['10.5'],'NA')
        self.assertEqual(results['10.4'],'Human Check')
        self.assertTrue(all(issue['addresses'][0].startswith('https://audit-client.test/') for issue in prepared['issues'] if issue.get('kind')=='human_check'))
        self.assertEqual(complete_for_export(prepared,'https://audit-client.test/'),prepared)
        with tempfile.TemporaryDirectory(dir=Path.cwd()) as directory:
            output=Path(directory)/'audit.xlsx'
            self.assertEqual(build(data,output,site_url='https://audit-client.test/')['checks'],7)
            wb=load_workbook(output); self.assertEqual(wb['Checklist'].max_row,8); wb.close()

    def test_completion_preserves_no_and_downgrades_provisional_pass_with_gap(self):
        data=fixture()
        data['evidence_gaps']=[{'check':'10.1','missing':'Secondary origin unverified.','next_action':'Verify secondary origin.'},
                               {'check':'10.4','missing':'Resource relationships missing.','next_action':'Obtain resource inlinks.'},
                               {'check':'10.5','missing':'Untested live transaction flow.','next_action':'Obtain read-only flow evidence.'}]
        prepared=complete_for_export(data,'https://example.com/')
        rows,_=normalize(prepared)
        self.assertEqual(rows[0]['result'],'No')
        self.assertIn('Human Check',rows[0]['findings'])
        self.assertEqual(rows[1]['result'],'Human Check')
        self.assertIn('Earlier provisional result: Yes',rows[1]['findings'])
        self.assertEqual(rows[2]['result'],'Human Check')
        self.assertIn('Earlier provisional result: NA',rows[2]['findings'])

    def test_empty_progress_can_export_and_invalid_evidence_still_rejected(self):
        prepared=complete_for_export({},'https://example.com/')
        rows,issues=normalize(prepared)
        self.assertEqual([row['result'] for row in rows],['Human Check']*7)
        self.assertEqual(len(issues),7)
        for url in ['not a URL','https://user:secret@example.com/']:
            with self.assertRaises(ValueError): complete_for_export({},url)
        data=fixture(); data['issues']=[]
        with self.assertRaises(ValueError): complete_for_export(data,'https://example.com/')
        data=fixture(); data['checklist'][0]['id']='10.9'
        with self.assertRaises(ValueError): complete_for_export(data,'https://example.com/')

    def test_identical_issue_retains_all_check_associations(self):
        data=fixture(); data['checklist'][1]['result']='No'
        duplicate=copy.deepcopy(data['issues'][0]); duplicate['check_ids']=['10.4']
        data['issues'].append(duplicate)
        _,issues=normalize(data)
        self.assertEqual(len(issues),1)
        self.assertEqual(issues[0]['check_ids'],['10.1','10.4'])

    def test_filename(self):
        self.assertEqual(report_filename('Example','2026-10-05'),'Example_robots_audit_2026-10-05.xlsx')
        self.assertEqual(report_filename('品牌/香港','2026-10-05'),'品牌_香港_robots_audit_2026-10-05.xlsx')
        with self.assertRaises(ValueError): report_filename('Example','2026-02-30')

    def test_export_schema_reasons_urls_and_no_overwrite(self):
        data=fixture(); data['issues'].append(copy.deepcopy(data['issues'][0]))
        data['issues'][0]['addresses'].append(data['issues'][0]['addresses'][0])
        with tempfile.TemporaryDirectory(dir=Path.cwd()) as directory:
            output=Path(directory)/'audit.xlsx'; result=build(data,output)
            self.assertEqual(result['issues'],1)
            wb=load_workbook(output)
            self.assertEqual(wb.sheetnames,['Checklist','Issues'])
            self.assertEqual([c.value for c in wb['Checklist'][1]],['Check','Result','Findings','Coverage'])
            self.assertEqual([c.value for c in wb['Issues'][1]],['Issue','Issue Description','How to Fix','Address'])
            self.assertIn(data['checklist'][2]['na_reason'],wb['Checklist']['C4'].value)
            self.assertEqual(wb['Checklist']['B3'].value,'Yes')
            self.assertEqual(wb['Checklist'].data_validations.dataValidation[0].formula1,'"Yes,No,NA,Human Check"')
            self.assertEqual(wb['Issues']['D2'].value,data['issues'][0]['addresses'][0])
            self.assertEqual(wb['Issues']['B2'].data_type,'s')
            self.assertEqual(len(wb['Checklist'].data_validations.dataValidation),1)
            wb.close()
            with self.assertRaises(FileExistsError): build(data,output)

    def test_rejects_na_without_reason_and_unknown_results(self):
        data=fixture(); del data['checklist'][2]['na_reason']
        with self.assertRaises(ValueError): normalize(data)
        for value in ['', 'Needs Review', 'Incomplete', None]:
            data=fixture(); data['checklist'][1]['result']=value
            with self.assertRaises(ValueError): normalize(data)

    def test_rejects_false_yes_missing_issue_and_invalid_links(self):
        data=fixture(); data['checklist'][0]['result']='Yes'
        with self.assertRaises(ValueError): normalize(data)
        data=fixture(); data['issues']=[]
        with self.assertRaises(ValueError): normalize(data)
        data=fixture(); data['issues'][0]['check_ids']=['10.7']
        with self.assertRaises(ValueError): normalize(data)

    def test_missing_requested_rows_or_unlabelled_gaps_are_rejected(self):
        data=fixture(); data['checklist'].pop()
        with self.assertRaises(ValueError): normalize(data)
        data=fixture(); data['evidence_gaps']=[{'check':'10.4','missing':'resource relationships'}]
        with self.assertRaises(ValueError): normalize(data)
        data=fixture(); del data['requested_checks']
        with self.assertRaises(ValueError): normalize(data)

    def test_human_check_is_exported_with_gap_issue_and_action(self):
        data=fixture()
        row=data['checklist'][1]
        row.update(result='Human Check', human_check_reason='Resource relationship export is missing.')
        data['evidence_gaps']=[{'check':'10.4','missing':'Resource relationship export is missing.', 'next_action':'Export source page/resource pairs and verify their responses.'}]
        data['issues'].append({'check_ids':['10.4'], 'kind':'human_check', 'issue':'Resource coverage unverified',
                              'description':'Required resources have not all been verified; no confirmed resource defect.',
                              'how_to_fix':'Export source page/resource pairs and verify their responses.', 'addresses':['https://example.com/']})
        with tempfile.TemporaryDirectory(dir=Path.cwd()) as directory:
            build(data,Path(directory)/'audit.xlsx')
            wb=load_workbook(Path(directory)/'audit.xlsx')
            self.assertEqual(wb['Checklist']['B3'].value,'Human Check')
            self.assertIn('Human Check',wb['Checklist']['C3'].value)
            self.assertIn('Export source page/resource pairs',wb['Checklist']['C3'].value)
            self.assertTrue(wb['Issues']['A3'].value.startswith('Human Check'))
            self.assertEqual(wb['Checklist']['B3'].fill.fgColor.rgb,'00FFF2CC')
            wb.close()

    def test_human_check_cannot_replace_confirmed_defect_or_link_to_yes(self):
        data=fixture()
        data['issues'][0]['kind']='human_check'
        with self.assertRaises(ValueError): normalize(data)
        data=fixture()
        data['issues'].append({'check_ids':['10.4'], 'kind':'human_check', 'issue':'Unknown', 'description':'Unknown',
                              'how_to_fix':'Verify', 'addresses':['https://example.com/']})
        with self.assertRaises(ValueError): normalize(data)
        data=fixture(); data['checklist'][1]['result']='Human Check'
        with self.assertRaises(ValueError): normalize(data)

    def test_all_seven_unresolved_checks_still_export_final_workbook(self):
        ids=[f'10.{i}' for i in range(1,8)]
        data={'checklist':[{'id':cid,'check':cid,'result':'Human Check','findings':'No usable evidence supplied.',
                           'human_check_reason':'Required evidence unavailable.','coverage':'Not verified.'} for cid in ids],
              'issues':[{'check_ids':[cid],'kind':'human_check','issue':'Evidence missing '+cid,
                         'description':'No confirmed defect; evidence missing.','how_to_fix':'Obtain required evidence and verify '+cid,
                         'addresses':['https://example.com/']} for cid in ids]}
        with tempfile.TemporaryDirectory(dir=Path.cwd()) as directory:
            result=build(data,Path(directory)/'audit.xlsx')
            self.assertEqual(result['checks'],7)
            wb=load_workbook(Path(directory)/'audit.xlsx')
            self.assertEqual([wb['Checklist'].cell(i,2).value for i in range(2,9)],['Human Check']*7)
            self.assertEqual(wb['Issues'].max_row,8)
            wb.close()

    def test_confirmed_no_with_review_issue_keeps_defect_and_visible_gap(self):
        data=fixture()
        data['issues'].append({'check_ids':['10.1'],'kind':'human_check','issue':'Other origin unverified',
                              'description':'Secondary origin response unavailable.','how_to_fix':'Check secondary origin robots response.',
                              'addresses':['https://www.example.com/robots.txt']})
        rows,issues=normalize(data)
        self.assertEqual(rows[0]['result'],'No')
        self.assertIn('Human Check',rows[0]['findings'])
        self.assertIn('Check secondary origin',rows[0]['findings'])
        self.assertEqual(len(issues),2)

    def test_duplicate_check_and_legacy_schema(self):
        data=fixture(); data['checklist'].append(copy.deepcopy(data['checklist'][0]))
        with self.assertRaises(ValueError): normalize(data)
        with self.assertRaises(ValueError): normalize({'overview':[]})

    def test_empty_issue_sheet_and_shared_resource_addresses(self):
        data={'requested_checks':['10.4'],'checklist':[{'id':'10.4','check':'10.4 CSS/JS','result':'Yes','findings':'None detected in complete export.','coverage':'2 synthetic pages.'}],'issues':[]}
        with tempfile.TemporaryDirectory(dir=Path.cwd()) as directory:
            build(data,Path(directory)/'audit.xlsx')
            wb=load_workbook(Path(directory)/'audit.xlsx'); self.assertEqual(wb['Issues'].max_row,1); wb.close()
        data=fixture(); data['issues'][0]['addresses']=['https://example.com/a','https://example.com/b','http://example.com/logo.png']
        _,issues=normalize(data)
        self.assertEqual(issues[0]['address'].splitlines(),data['issues'][0]['addresses'])


if __name__=='__main__': unittest.main()
