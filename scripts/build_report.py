"""Render agent-reviewed robots.txt findings. No network or MCP calls."""
import argparse
import copy
import json
import math
import re
from datetime import date
from pathlib import Path
from urllib.parse import urlsplit

CHECK_IDS = {f"10.{i}" for i in range(1, 8)}
CHECK_NAMES = {
    '10.1': 'Does a robots.txt document exist?',
    '10.2': 'Are sections of the site that we want indexed, being disallowed with robots.txt?',
    '10.3': "Are disallow blocks in place for URLs that Google shouldn't be crawling?",
    '10.4': 'Are we allowing bot access to CSS and JS documents?',
    '10.5': 'Are the following pages disallow? shopping carts, thank you, admin, duplicate content, account pages.',
    '10.6': 'Is the robots.txt called robots.txt, and not ROBOTS.TXT?',
    '10.7': 'Is there a link in the robots.txt file to the sitemap index/sitemap?',
}

REVIEW_ACTIONS = {
    '10.1': 'Open /robots.txt and check whether it loads and contains robots rules or a sitemap link.',
    '10.2': 'Check whether robots.txt allows Google to crawl the pages still awaiting review.',
    '10.3': 'Identify which pages should be kept out of crawling, then check whether robots.txt blocks them.',
    '10.4': 'Export the pages that use these files and their crawl results.',
    '10.5': 'Confirm what these pages do and how access is restricted.',
    '10.6': 'Check whether the lowercase /robots.txt address loads the correct file.',
    '10.7': 'Check whether robots.txt contains a complete sitemap link, starting with http:// or https://.',
}

SHEETS = {
    "Checklist": (["Item No.", "Item Name", "Initial Check", "Findings", "Instruction"], ["id", "item_name", "result", "findings", "instruction"]),
    "10. Robot.txt": (["Issue", "Issue Description", "How to Fix", "Address"], ["issue", "description", "how_to_fix", "address"]),
}


def report_filename(site_name, audit_date, check_ids=None):
    if not isinstance(site_name, str) or not site_name.strip():
        raise ValueError("site name is required")
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", audit_date):
        raise ValueError("audit date must use YYYY-MM-DD")
    date.fromisoformat(audit_date)
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", site_name.strip()).strip(" .")
    if not name:
        raise ValueError("site name contains no usable filename characters")
    ids = list(CHECK_IDS) if check_ids is None else list(check_ids)
    if not ids or any(cid not in CHECK_IDS for cid in ids):
        raise ValueError('Filename requires valid covered robots check IDs')
    ids = sorted(set(ids), key=lambda cid: tuple(map(int, cid.split('.'))))
    covered = ids[0] if len(ids) == 1 else f'{ids[0]}-{ids[-1]}'
    return f"{covered}_{name}_robots_audit_{audit_date}.xlsx"


def required_text(row, fields):
    for field in fields:
        if not isinstance(row.get(field), str) or not row[field].strip():
            raise ValueError(f"Missing/non-string {field}")
        if len(row[field]) > 32767:
            raise ValueError(f"{field} exceeds Excel cell limit; split issue groups")


def add_instruction(row, text):
    """Keep human actions separate and retain each exact action once."""
    if not isinstance(text, str):
        raise ValueError('instruction must be a string')
    lines = [] if row.get('instruction', '') in {'', '—'} else row['instruction'].splitlines()
    for line in text.splitlines():
        line = line.strip()
        if line and line != '—' and line not in lines:
            lines.append(line)
    row['instruction'] = '\n'.join(lines)


def separate_legacy_actions(row):
    """Migrate explicitly labelled old actions; never guess prose meaning."""
    pattern = r'(?im)(?:^|\n|[ \t]+)(?:Next action|下一步(?:动作)?|Instruction)\s*[:：]\s*([^\n]*)'
    def move(match):
        add_instruction(row, match.group(1))
        return ''
    row['findings'] = re.sub(pattern, move, row['findings']).strip()


def numbered_issue(row):
    ids = sorted(set(row['check_ids']), key=lambda cid: tuple(map(int, cid.split('.'))))
    # Derive numbers from associations, including already-prefixed legacy titles.
    title = re.sub(r'^10\.[1-7](?:\s*[,/、]\s*10\.[1-7])*[ :：-]+', '', row['issue'])
    return ', '.join(ids) + ' ' + title


def rich_findings(text, phrases=()):
    """Apply Excel rich-text emphasis without changing the Findings text."""
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont

    labels = r'(?:Human Check|Already checked(?: / available evidence)?|Missing(?:/failed)?|Next action|Conclusion|Excluded(?: resources)?|Coverage|已检查|已验证|检查结果|结论|缺少(?:／失败)?|缺口(?:／错误)?|失败|下一步(?:动作)?|排除原因|已排除|覆盖范围|匹配规则)'
    intervals = [(match.start(1), match.end(1)) for match in re.finditer(r'^\s*(' + labels + r'\s*[:：])', text, re.M | re.I)]
    for phrase in phrases:
        intervals.extend((match.start(), match.end()) for match in re.finditer(re.escape(phrase), text))
    if not intervals:
        # Unstructured findings still get a short leading conclusion emphasized.
        match = re.search(r'\S[^\n]*', text)
        if match and not match.group().startswith(('http://', 'https://')):
            leading = match.group()
            sentence_end = re.search(r'[。！？]|[.!?](?=\s|$)', leading)
            end = sentence_end.end() if sentence_end else len(leading)
            if end <= 160:
                intervals.append((match.start(), match.start() + end))
    merged = []
    for start, end in sorted(intervals):
        if merged and start <= merged[-1][1]:
            merged[-1] = (merged[-1][0], max(end, merged[-1][1]))
        else:
            merged.append((start, end))
    if not merged:
        return text
    runs, cursor = [], 0
    for start, end in merged:
        if start > cursor:
            runs.append(text[cursor:start])
        runs.append(TextBlock(InlineFont(rFont='Calibri', sz=11, b=True), text[start:end]))
        cursor = end
    if cursor < len(text):
        runs.append(text[cursor:])
    return CellRichText(runs)


def complete_for_export(data, site_url):
    """Explicit delivery fallback, using the actual audited origin; never invent a pass.

    Strict normalize remains the final validator. Invalid IDs, unsupported results,
    and unsubstantiated No/defect associations are not silently repaired.
    """
    if not isinstance(data, dict) or not isinstance(site_url, str):
        raise ValueError('Completion requires a report object and actual site URL string')
    parsed = urlsplit(site_url)
    if parsed.scheme not in {'http', 'https'} or not parsed.hostname or parsed.username or parsed.password or any(c.isspace() for c in site_url):
        raise ValueError('site_url must be the actual absolute HTTP(S) audited URL without credentials')
    result = copy.deepcopy(data)
    if 'overview' in result:
        raise ValueError('Legacy schema must be migrated before completion')
    requested = result.get('requested_checks', sorted(CHECK_IDS))
    if not isinstance(requested, list) or not requested or any(not isinstance(cid, str) or cid not in CHECK_IDS for cid in requested) or len(set(requested)) != len(requested):
        raise ValueError('requested_checks must contain unique supported check IDs')
    rows = result.setdefault('checklist', [])
    issues = result.setdefault('issues', [])
    gaps = result.setdefault('evidence_gaps', [])
    if not all(isinstance(value, list) for value in (rows, issues, gaps)):
        raise ValueError('checklist, issues and evidence_gaps must be lists')
    ids = {}
    for row in rows:
        if not isinstance(row, dict) or row.get('id') not in requested or row['id'] in ids:
            raise ValueError('Unknown/duplicate/unrequested check ID')
        ids[row['id']] = row
    by_check = {}
    for gap in gaps:
        if not isinstance(gap, dict):
            raise ValueError('Each evidence gap must be an object')
        required_text(gap, ['check', 'missing', 'next_action'])
        if gap['check'] not in requested:
            raise ValueError('Evidence gap references an unrequested check')
        by_check.setdefault(gap['check'], []).append(gap)
    for cid in requested:
        if cid not in ids:
            row = {'id': cid, 'check': cid + ' ' + CHECK_NAMES[cid], 'result': 'Human Check',
                   'findings': 'There is not enough information to confirm this result.',
                   'coverage': 'No completed checks were recorded for this item.'}
            rows.append(row)
            ids[cid] = row
        row = ids[cid]
        row['result'] = {'√': 'Yes', 'X': 'No'}.get(row.get('result'), row.get('result'))
        check_gaps = by_check.get(cid, [])
        if check_gaps and row.get('result') in {'Yes', 'NA'}:
            row['findings'] = 'Earlier provisional result: ' + row['result'] + '.\n' + row.get('findings', '')
            row['result'] = 'Human Check'
        if row.get('result') != 'Human Check' and not (row.get('result') == 'No' and check_gaps):
            continue
        reason = '\n'.join(gap['missing'] for gap in check_gaps) or row.get('human_check_reason') or 'There is not enough information to confirm this result.'
        if row['result'] == 'Human Check':
            row.setdefault('human_check_reason', reason)
            if not row['human_check_reason']:
                row['human_check_reason'] = reason
        action = '\n'.join(gap['next_action'] for gap in check_gaps) or row.get('instruction') or row.get('human_check_action') or REVIEW_ACTIONS[cid]
        row['human_check_action'] = action
    completed_rows, defects = normalize(result)
    result['checklist'] = completed_rows
    result['issues'] = defects
    return result


def normalize(data):
    if "overview" in data:
        raise ValueError("Legacy schema: migrate to checklist/issues with separate description and fix")
    gaps = data.get("evidence_gaps", [])
    if not isinstance(gaps, list):
        raise ValueError("evidence_gaps must be a list")
    requested = data.get("requested_checks", sorted(CHECK_IDS))
    if not isinstance(requested, list) or not requested or any(not isinstance(cid, str) or cid not in CHECK_IDS for cid in requested) or len(set(requested)) != len(requested):
        raise ValueError("requested_checks must contain unique supported check IDs")
    checks, issues = data.get("checklist"), data.get("issues")
    if not isinstance(checks, list) or not checks or not isinstance(issues, list):
        raise ValueError("Non-empty checklist and issues list required")
    ids = {}
    for source in checks:
        row = dict(source)
        required_text(row, ["id", "check", "findings"])
        row['result'] = {'√': 'Yes', 'X': 'No'}.get(row.get('result'), row.get('result'))
        row.setdefault('coverage', '')  # Optional archived evidence, not an Excel column.
        row.setdefault('instruction', '')
        if not isinstance(row['instruction'], str):
            raise ValueError('instruction must be a string')
        separate_legacy_actions(row)
        phrases = row.get('findings_bold', [])
        if not isinstance(phrases, list) or any(not isinstance(phrase, str) or not phrase.strip() for phrase in phrases):
            raise ValueError('findings_bold must be a list of non-empty exact text phrases')
        if row["id"] not in CHECK_IDS or row["id"] in ids:
            raise ValueError("Unknown/duplicate check ID")
        # Checklist titles are fixed user-supplied text, independent of report language.
        row["check"] = row["id"] + " " + CHECK_NAMES[row["id"]]
        if row.get("result") not in {"Yes", "No", "NA", "Human Check"}:
            raise ValueError("Final Result must be Yes, No, NA or Human Check")
        reason_field = {"NA": "na_reason", "Human Check": "human_check_reason"}.get(row["result"])
        if reason_field:
            required_text(row, [reason_field])
            if row[reason_field] not in row["findings"]:
                row["findings"] += "\n" + row[reason_field]
        required_text(row, ["findings"])
        if row["result"] == "Human Check" and "Human Check" not in row["findings"]:
            row["findings"] = "Human Check: " + row["findings"]
        required_text(row, ["findings"])
        ids[row["id"]] = row
    if set(ids) != set(requested):
        raise ValueError("Final report must include every requested check; do not silently omit unresolved checks")
    unique, linked = {}, set()
    for source in issues:
        row = dict(source)
        required_text(row, ["issue", "description", "how_to_fix"])
        refs = row.get("check_ids")
        kind = row.get("kind", "defect")
        if kind not in {"defect", "human_check"}:
            raise ValueError("Issue kind must be defect or human_check")
        permitted = {"No"} if kind == "defect" else {"No", "Human Check"}
        if not isinstance(refs, list) or not refs or any(not isinstance(cid, str) or cid not in ids or ids[cid]["result"] not in permitted for cid in refs):
            raise ValueError("Defects link to No; human_check issues link to No or Human Check")
        if kind == "human_check":
            # Compatibility: move legacy review rows to Checklist, never the defect tab.
            marker = 'Human Check: ' + row['issue'] + '\n' + row['description']
            addresses = row.get('addresses', [])
            if not isinstance(addresses, list) or any(not isinstance(url, str) for url in addresses):
                raise ValueError('Legacy review addresses must be a list of strings')
            for cid in refs:
                add_instruction(ids[cid], row['how_to_fix'])
                if marker not in ids[cid]['findings']:
                    ids[cid]['findings'] += '\n' + marker
                if addresses:
                    coverage_note = 'Review addresses: ' + '\n'.join(dict.fromkeys(addresses))
                    if coverage_note not in ids[cid]['coverage']:
                        ids[cid]['coverage'] += '\n' + coverage_note
                required_text(ids[cid], ['findings', 'instruction'])
            continue
        linked.update(refs)
        required_text(row, ["issue"])
        addresses = row.get("addresses")
        if not isinstance(addresses, list) or not addresses or any(not isinstance(url, str) or not re.match(r"^https?://[^\s/]+(?:/[^\s]*)?$", url) for url in addresses):
            raise ValueError("Issue addresses must be absolute HTTP(S) URLs")
        row["address"] = "\n".join(dict.fromkeys(addresses))
        required_text(row, ["address"])
        identity = (kind,) + tuple(row[field] for field in ["issue", "description", "how_to_fix", "address"])
        if identity in unique:
            unique[identity]['check_ids'] = list(dict.fromkeys(unique[identity]['check_ids'] + refs))
        else:
            unique[identity] = row
    if {cid for cid, row in ids.items() if row["result"] == "No"} != linked:
        raise ValueError("Every No requires a confirmed-defect row in 10. Robot.txt")
    # Keep the repair step even when a separate evidence gap also needs follow-up.
    for cid, row in ids.items():
        if not row['instruction'] and row['result'] == 'No':
            for issue in unique.values():
                if cid in issue['check_ids']:
                    add_instruction(row, issue['how_to_fix'])
    for gap in gaps:
        if not isinstance(gap, dict):
            raise ValueError("Each evidence gap must be an object")
        required_text(gap, ["check", "missing", "next_action"])
        cid = gap["check"]
        if cid not in ids or ids[cid]["result"] not in {"No", "Human Check"}:
            raise ValueError("Every evidence gap must reference a No or Human Check Checklist row")
        if gap["missing"] not in ids[cid]["findings"]:
            ids[cid]["findings"] += "\nHuman Check: " + gap["missing"]
        add_instruction(ids[cid], gap['next_action'])
        required_text(ids[cid], ["findings"])
    for cid, row in ids.items():
        # A supplied human instruction wins over generic default wording.
        if row.get('human_check_action'):
            add_instruction(row, row['human_check_action'])
        if not row['instruction'] and row['result'] == 'Human Check':
            add_instruction(row, REVIEW_ACTIONS[cid])
        if not row['instruction']:
            row['instruction'] = '—'
        required_text(row, ['findings', 'instruction'])
    return sorted(ids.values(), key=lambda r: tuple(map(int, r["id"].split(".")))), list(unique.values())


def build(data, output, site_url=None):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    if site_url is not None:
        data = complete_for_export(data, site_url)
    checklist, issues = normalize(data)
    output = Path(output)
    if output.exists():
        raise FileExistsError(f"Refusing to overwrite {output}; select a new path")
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    # Keep internal legacy Yes/No compatibility; Excel only displays the symbols.
    checklist = [dict(row, item_name=CHECK_NAMES[row['id']], result={'Yes': '√', 'No': 'X'}.get(row['result'], row['result'])) for row in checklist]
    issues = [dict(row, issue=numbered_issue(row)) for row in issues]
    datasets = {"Checklist": checklist, "10. Robot.txt": issues}
    widths = {"Checklist": [12, 64, 16, 90, 64], "10. Robot.txt": [42, 90, 80, 90]}
    for title, (headers, fields) in SHEETS.items():
        sheet = wb.create_sheet(title)
        sheet.append(headers)
        for row in datasets[title]:
            sheet.append([row[field] for field in fields])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 25
        for cell in sheet[1]:
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F75B5")
        for j, width in enumerate(widths[title], 1):
            sheet.column_dimensions[sheet.cell(1, j).column_letter].width = width
        for row_index, excel_row in enumerate(sheet.iter_rows(min_row=2)):
            line_count = 1
            for j, cell in enumerate(excel_row):
                cell.font = Font(name="Calibri", size=11)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                    cell.number_format = "@"
                    line_count = max(line_count, sum(max(1, math.ceil(len(part) / max(1, widths[title][j] // 2))) for part in cell.value.split("\n")))
                    if "\n" not in cell.value and cell.value.startswith(("https://", "http://")):
                        cell.hyperlink = cell.value
                        cell.font = Font(name="Calibri", size=11, color="0563C1", underline="single")
            sheet.row_dimensions[excel_row[0].row].height = min(409, max(30, line_count * 16 + 8))
            if title == "Checklist":
                source = checklist[row_index]
                excel_row[3].value = rich_findings(source['findings'], source.get('findings_bold', []))
                excel_row[3].data_type = 's'
                color = {"X": "FCE4D6", "NA": "E7E6E6", "Human Check": "FFF2CC"}.get(excel_row[2].value)
                if color:
                    excel_row[2].fill = PatternFill("solid", fgColor=color)
        if title == "Checklist":
            validation = DataValidation(type="list", formula1='"√,X,NA,Human Check"', allow_blank=False)
            validation.showErrorMessage = True
            sheet.add_data_validation(validation)
            validation.add(f"C2:C{sheet.max_row}")
    wb.save(output)
    verified = load_workbook(output)
    try:
        if verified.sheetnames != list(SHEETS):
            raise RuntimeError("Workbook sheet verification failed")
        for title, (headers, fields) in SHEETS.items():
            sheet = verified[title]
            if [c.value for c in sheet[1]] != headers or sheet.max_row != len(datasets[title]) + 1:
                raise RuntimeError(f"Workbook header/count verification failed: {title}")
            for i, source in enumerate(datasets[title], 2):
                if [sheet.cell(i, j).value or "" for j in range(1, len(fields) + 1)] != [source[f] for f in fields]:
                    raise RuntimeError("Workbook value preservation failed")
    finally:
        verified.close()
    return {"output": str(output.resolve()), "checks": len(checklist), "issues": len(issues)}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--site-name", required=True)
    parser.add_argument("--date", required=True, help="YYYY-MM-DD in the user's timezone")
    parser.add_argument("--site-url", help="Actual audited URL: explicitly complete missing checks/review details in Checklist as Human Check")
    parser.add_argument("--prepared-input", type=Path, help="Archive completed report input to a new JSON path (requires --site-url)")
    args = parser.parse_args()
    data = json.loads(args.input.read_text(encoding="utf-8-sig"))
    if args.prepared_input and not args.site_url:
        parser.error('--prepared-input requires --site-url')
    if args.site_url:
        data = complete_for_export(data, args.site_url)
    checklist, _ = normalize(data)
    output = args.output_dir / report_filename(args.site_name, args.date, [row['id'] for row in checklist])
    if output.exists():
        raise FileExistsError(f'Refusing to overwrite {output}; select a new path')
    if args.prepared_input:
        args.prepared_input.parent.mkdir(parents=True, exist_ok=True)
        with args.prepared_input.open('x', encoding='utf-8') as stream:
            stream.write(json.dumps(data, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(build(data, output), ensure_ascii=False))
